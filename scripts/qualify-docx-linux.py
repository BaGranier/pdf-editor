"""DOCX editability, accents, simple tables and images, reopened by LibreOffice."""

from __future__ import annotations
import importlib
import json
import sys
from pathlib import Path

import fitz
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "services/pdf-engine"))
from app.conversion.docx_converter import PdfToDocxConverter  # noqa: E402
from app.conversion.models import DocxMode  # noqa: E402

render = importlib.import_module("scripts.validate-docx-visual-quality")
OUTPUT = ROOT / "data/output/editor-robustness-008/docx"


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    results = []
    for fixture, witness in [
        ("conversion-simple-text.pdf", "été"),
        ("conversion-table.pdf", "Quantité"),
        ("conversion-images.pdf", "image intégrée"),
    ]:
        source = ROOT / "apps/web/e2e/fixtures" / fixture
        docx = OUTPUT / (source.stem + ".docx")
        PdfToDocxConverter().convert(source, docx, mode=DocxMode.EDITABLE)
        document = Document(docx)
        text = "\n".join(
            [
                *(p.text for p in document.paragraphs),
                *(
                    cell.text
                    for table in document.tables
                    for row in table.rows
                    for cell in row.cells
                ),
            ]
        )
        assert witness in text
        if "table" in fixture:
            assert document.tables
        if "images" in fixture:
            assert any("image" in rel.reltype for rel in document.part.rels.values())
        pdf = render.convert_with_libreoffice(docx, OUTPUT)
        with fitz.open(pdf) as reopened, fitz.open(source) as original:
            reopened_text = " ".join(page.get_text() for page in reopened)
            assert witness in reopened_text
            if "images" in fixture:
                assert any(page.get_images() for page in reopened)
            results.append(
                dict(
                    fixture=fixture,
                    sourcePages=len(original),
                    renderedPages=len(reopened),
                    editableParagraphs=len(document.paragraphs),
                    tables=len(document.tables),
                    imageRelationships=sum(
                        "image" in rel.reltype for rel in document.part.rels.values()
                    ),
                    witness=True,
                )
            )
    (OUTPUT / "results.json").write_text(json.dumps(results, indent=2))
    print(json.dumps(results, indent=2))


if __name__ == "__main__":
    main()
