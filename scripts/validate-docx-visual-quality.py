#!/usr/bin/env python3
"""Validate DOCX fidelity with a synthetic PDF and LibreOffice rendering."""

from __future__ import annotations

import argparse
import json
import math
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any

import fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml.ns import qn


PROJECT_ROOT = Path(__file__).resolve().parents[1]
BACKEND_DIR = PROJECT_ROOT / "services" / "pdf-engine"
SOURCE_PDF = (
    PROJECT_ROOT
    / "apps"
    / "web"
    / "e2e"
    / "fixtures"
    / "conversion-docx-fidelity.pdf"
)
OUTPUT_DIR = (
    PROJECT_ROOT / "apps" / "web" / "test-results" / "docx-visual-quality"
)

sys.path.insert(0, str(BACKEND_DIR))

from app.conversion.docx_converter import PdfToDocxConverter  # noqa: E402
from app.conversion.models import DocxMode  # noqa: E402


def page_ink_ratio(page: fitz.Page) -> float:
    pixmap = page.get_pixmap(matrix=fitz.Matrix(0.35, 0.35), alpha=False)
    samples = pixmap.samples
    dark = 0
    for index in range(0, len(samples), pixmap.n):
        if all(samples[index + channel] < 245 for channel in range(3)):
            dark += 1
    return dark / (pixmap.width * pixmap.height)


def black_pixel_ratio(image: bytes) -> float:
    pixmap = fitz.Pixmap(image)
    samples = pixmap.samples
    dark = 0
    channels = min(3, pixmap.n)
    for index in range(0, len(samples), pixmap.n):
        opaque = pixmap.n < 4 or samples[index + pixmap.n - 1] >= 32
        if opaque and all(
            samples[index + channel] < 24
            for channel in range(channels)
        ):
            dark += 1
    return dark / (pixmap.width * pixmap.height)


def image_non_white_ratio(image: bytes) -> float:
    pixmap = fitz.Pixmap(image)
    samples = pixmap.samples
    non_white = 0
    for index in range(0, len(samples), pixmap.n):
        opaque = pixmap.n < 4 or samples[index + pixmap.n - 1] >= 32
        if opaque and any(
            samples[index + channel] < 248
            for channel in range(min(3, pixmap.n))
        ):
            non_white += 1
    return round(non_white / (pixmap.width * pixmap.height), 5)


def render_comparisons(
    source: fitz.Document,
    rendered: fitz.Document,
    output_directory: Path,
    prefix: str,
) -> list[str]:
    captures: list[str] = []
    comparison_count = max(len(source), len(rendered))
    for page_index in range(comparison_count):
        source_page = source[min(page_index, len(source) - 1)]
        rendered_page = rendered[min(page_index, len(rendered) - 1)]
        width = max(source_page.rect.width, rendered_page.rect.width)
        height = max(source_page.rect.height, rendered_page.rect.height)
        comparison = fitz.open()
        page = comparison.new_page(width=width * 2 + 18, height=height)
        page.show_pdf_page(
            fitz.Rect(0, 0, width, height),
            source,
            min(page_index, len(source) - 1),
            keep_proportion=True,
        )
        page.show_pdf_page(
            fitz.Rect(width + 18, 0, width * 2 + 18, height),
            rendered,
            min(page_index, len(rendered) - 1),
            keep_proportion=True,
        )
        capture = (
            output_directory
            / f"comparison-{prefix}-page-{page_index + 1:02d}.png"
        )
        page.get_pixmap(dpi=110, alpha=False).save(capture)
        comparison.close()
        captures.append(str(capture.relative_to(PROJECT_ROOT)))
    return captures


def normalized_text_displacement(
    source_page: fitz.Page,
    rendered_page: fitz.Page,
    witness: str,
) -> float | None:
    source_matches = source_page.search_for(witness)
    rendered_matches = rendered_page.search_for(witness)
    if not source_matches or not rendered_matches:
        return None
    source_center = (
        source_matches[0].x0 + source_matches[0].x1
    ) / (2 * source_page.rect.width), (
        source_matches[0].y0 + source_matches[0].y1
    ) / (2 * source_page.rect.height)
    rendered_center = (
        rendered_matches[0].x0 + rendered_matches[0].x1
    ) / (2 * rendered_page.rect.width), (
        rendered_matches[0].y0 + rendered_matches[0].y1
    ) / (2 * rendered_page.rect.height)
    return round(
        math.hypot(
            source_center[0] - rendered_center[0],
            source_center[1] - rendered_center[1],
        ),
        4,
    )


def convert_with_libreoffice(
    docx_path: Path,
    output_directory: Path,
) -> Path:
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if executable is None:
        raise FileNotFoundError("LibreOffice headless n'est pas installé.")
    profile_directory = Path(
        tempfile.mkdtemp(prefix="docx-quality-lo-", dir=output_directory)
    )
    cache_directory = profile_directory / "cache"
    config_directory = profile_directory / "config"
    cache_directory.mkdir()
    config_directory.mkdir()
    environment = {
        **os.environ,
        "XDG_CACHE_HOME": str(cache_directory),
        "XDG_CONFIG_HOME": str(config_directory),
    }
    profile_uri = (profile_directory / "profile").resolve().as_uri()
    result = subprocess.run(
        [
            executable,
            f"-env:UserInstallation={profile_uri}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_directory),
            str(docx_path),
        ],
        check=False,
        capture_output=True,
        text=True,
        timeout=120,
        env=environment,
    )
    output_pdf = output_directory / f"{docx_path.stem}.pdf"
    shutil.rmtree(profile_directory, ignore_errors=True)
    if result.returncode != 0 or not output_pdf.is_file():
        diagnostic = (result.stderr or result.stdout).strip()
        raise RuntimeError(
            "LibreOffice n'a pas produit le PDF attendu."
            + (f" {diagnostic[-800:]}" if diagnostic else "")
        )
    return output_pdf


def inspect_editable_docx(path: Path) -> dict[str, Any]:
    document = Document(path)
    title = next(
        (
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text == "Engagement individuel"
        ),
        None,
    )
    image_widths = [round(shape.width.pt, 2) for shape in document.inline_shapes]
    media_parts = [
        part
        for part in document.part.package.parts
        if part.content_type.startswith("image/")
        and "thumbnail" not in str(part.partname)
    ]
    return {
        "imageCount": len(media_parts),
        "inlineImageCount": len(document.inline_shapes),
        "imageWidthsPt": image_widths,
        "imageSizeRatios": [
            round(width / 120, 3) for width in image_widths
        ],
        "blackPixelRatio": (
            round(black_pixel_ratio(media_parts[0].blob), 5)
            if media_parts
            else None
        ),
        "witnessTextPresent": title is not None,
        "titleCentered": (
            title is not None
            and title.alignment == WD_ALIGN_PARAGRAPH.CENTER
        ),
        "boldPresent": any(
            run.bold
            for paragraph in document.paragraphs
            for run in paragraph.runs
            if "respect de ces règles" in paragraph.text
        ),
        "highlightPresent": any(
            run.font.highlight_color == WD_COLOR_INDEX.YELLOW
            for paragraph in document.paragraphs
            for run in paragraph.runs
            if "Nom de l'étudiant" in paragraph.text
        ),
        "borderPresent": any(
            paragraph._p.xpath(".//w:pBdr")
            for paragraph in document.paragraphs
            if "premier paragraphe encadré" in paragraph.text
        ),
        "listPresent": any(
            paragraph.style.name in {"List Bullet", "List Number"}
            for paragraph in document.paragraphs
        ),
    }


def inspect_visual_docx(path: Path) -> dict[str, Any]:
    document = Document(path)
    media_parts = [
        part
        for part in document.part.package.parts
        if part.content_type.startswith("image/")
        and "thumbnail" not in str(part.partname)
    ]
    image_paragraphs = [
        paragraph
        for paragraph in document.paragraphs
        if paragraph._p.xpath(".//w:drawing")
    ]
    exact_line_rule_count = 0
    for paragraph in image_paragraphs:
        spacing_nodes = paragraph._p.xpath("./w:pPr/w:spacing")
        if spacing_nodes and spacing_nodes[0].get(qn("w:lineRule")) == "exact":
            exact_line_rule_count += 1
    extent_ratios = []
    for index, shape in enumerate(document.inline_shapes):
        section = document.sections[min(index, len(document.sections) - 1)]
        extent_ratios.append(
            {
                "width": round(shape.width / section.page_width, 4),
                "height": round(shape.height / section.page_height, 4),
            }
        )
    return {
        "visualImageCount": len(document.inline_shapes),
        "visualImageParagraphCount": len(image_paragraphs),
        "visualImageNonWhiteRatios": [
            image_non_white_ratio(part.blob) for part in media_parts
        ],
        "visualImageExtentRatios": extent_ratios,
        "visualExactLineRuleParagraphs": exact_line_rule_count,
        "visualClippingDetected": exact_line_rule_count > 0,
    }


def validate(output_directory: Path) -> dict[str, Any]:
    output_directory.mkdir(parents=True, exist_ok=True)
    for generated_path in (
        output_directory / "editable.docx",
        output_directory / "editable.pdf",
        output_directory / "visual.docx",
        output_directory / "visual.pdf",
        *output_directory.glob("comparison-*-page-*.png"),
    ):
        generated_path.unlink(missing_ok=True)
    editable_docx = output_directory / "editable.docx"
    visual_docx = output_directory / "visual.docx"
    converter = PdfToDocxConverter()
    converter.convert(SOURCE_PDF, editable_docx, mode=DocxMode.EDITABLE)
    converter.convert(SOURCE_PDF, visual_docx, mode=DocxMode.VISUAL)

    editable_structure = inspect_editable_docx(editable_docx)
    visual_structure = inspect_visual_docx(visual_docx)
    source = fitz.open(SOURCE_PDF)
    source_pages = len(source)
    source_images = sum(len(page.get_images(full=True)) for page in source)
    try:
        editable_pdf = convert_with_libreoffice(editable_docx, output_directory)
        visual_pdf = convert_with_libreoffice(visual_docx, output_directory)
    except (FileNotFoundError, RuntimeError, subprocess.TimeoutExpired) as error:
        source.close()
        return {
            "status": "unavailable",
            "reason": str(error),
            "sourcePageCount": source_pages,
            "sourceImageCount": source_images,
            **editable_structure,
            **visual_structure,
            "captures": [],
        }

    editable_render = fitz.open(editable_pdf)
    visual_render = fitz.open(visual_pdf)
    editable_text = "\n".join(page.get_text() for page in editable_render)
    text_displacement = normalized_text_displacement(
        source[0],
        editable_render[0],
        "Engagement individuel",
    )
    checks = {
        "paginationReasonable": len(editable_render) <= source_pages + 1,
        "visualPaginationExact": len(visual_render) == source_pages,
        "editableNoBlankPage": all(
            page_ink_ratio(page) > 0.002 for page in editable_render
        ),
        "visualNoBlankPage": all(
            page_ink_ratio(page) > 0.002 for page in visual_render
        ),
        "noBlackBackground": (
            editable_structure["blackPixelRatio"] is not None
            and editable_structure["blackPixelRatio"] < 0.05
        ),
        "imageSizeReasonable": all(
            0.8 <= ratio <= 1.2
            for ratio in editable_structure["imageSizeRatios"]
        ),
        "imageRetained": editable_structure["imageCount"] > 0,
        "textRetained": (
            editable_structure["witnessTextPresent"]
            and "Engagement individuel" in editable_text
        ),
        "textNotMassivelyMoved": (
            text_displacement is not None and text_displacement < 0.25
        ),
        "boldRetained": editable_structure["boldPresent"],
        "highlightRetained": editable_structure["highlightPresent"],
        "borderRetained": editable_structure["borderPresent"],
        "listRetained": editable_structure["listPresent"],
        "visualImagesHaveContent": all(
            ratio > 0.002
            for ratio in visual_structure["visualImageNonWhiteRatios"]
        ),
        "visualImageExtentsMatchPages": all(
            0.97 <= ratio["width"] <= 1
            and 0.97 <= ratio["height"] <= 1
            for ratio in visual_structure["visualImageExtentRatios"]
        ),
        "visualImagesNotClipped": not visual_structure[
            "visualClippingDetected"
        ],
    }
    captures = [
        *render_comparisons(
            source,
            editable_render,
            output_directory,
            "editable",
        ),
        *render_comparisons(
            source,
            visual_render,
            output_directory,
            "visual",
        ),
    ]
    payload = {
        "status": "passed" if all(checks.values()) else "failed",
        "sourcePageCount": source_pages,
        "renderedPageCount": {
            "editable": len(editable_render),
            "visual": len(visual_render),
        },
        "sourceImageCount": source_images,
        "textDisplacementRatio": text_displacement,
        **editable_structure,
        **visual_structure,
        "checks": checks,
        "captures": captures,
    }
    source.close()
    editable_render.close()
    visual_render.close()
    return payload


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--required",
        action="store_true",
        help="Échouer également si LibreOffice est indisponible.",
    )
    arguments = parser.parse_args()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    payload = validate(OUTPUT_DIR)
    results_path = OUTPUT_DIR / "results.json"
    results_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(f"Rapport qualité DOCX : {results_path}")
    if payload["status"] == "failed":
        return 1
    if payload["status"] == "unavailable" and arguments.required:
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
